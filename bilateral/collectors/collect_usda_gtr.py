#!/usr/bin/env python3
"""
collect_usda_gtr.py — AgriMacro Intelligence
Coleta dados do USDA AMS Grain Transportation Report (GTR)
Fonte: https://www.ams.usda.gov/services/transportation-analysis/gtr-datasets

Tabelas coletadas:
  Table 1  - Cost indicators: diesel ($/gal), rail ($/car), barge (% tariff), ocean Gulf & PNW ($/mt)
  Table 2  - Price spreads: origin to export position ($/bu) — basis proxy
  Table 7  - Rail tariffs: corn & soybean shuttle/unit trains ($/car + fuel surcharge)
  Table 9  - Barge grain movements: volumes by lock (tons)

Output: JSON padronizado AgriMacro em data/usda_gtr/
Uso bilateral: Lado EUA do "Landed Cost Spread"
Principio: ZERO MOCK — apenas dados reais do USDA AMS
"""

import json
import logging
import os
import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path

import requests
import urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl required. Run: pip install openpyxl")
    sys.exit(1)

# ─────────────────────────────────────────────
# CONFIG
# ─────────────────────────────────────────────

BASE_URL = "https://www.ams.usda.gov/sites/default/files/media"

TABLES = {
    "table1": {"url": f"{BASE_URL}/GTRTable1.xlsx", "sheet": "Data",
               "label": "Grain Transport Cost Indicators"},
    "table2": {"url": f"{BASE_URL}/GTRTable2.xlsx", "sheet": "Data",
               "label": "U.S. Origins to Export Price Spreads"},
    "table7": {"url": f"{BASE_URL}/GTRTable7.xlsx", "sheet": "GTR Table 7",
               "label": "Rail Tariff Rates - Corn & Soybean"},
    "table9": {"url": f"{BASE_URL}/GTRTable9.xlsx", "sheet": "Socrata_Data",
               "label": "Barge Grain Movements"},
}

BASE_DIR = Path(os.environ.get("AGRIMACRO_DATA_DIR", "data"))
OUTPUT_DIR = BASE_DIR / "usda_gtr"
CACHE_DIR = OUTPUT_DIR / "cache"
RAW_DIR = OUTPUT_DIR / "raw_xlsx"

LOG_FORMAT = "%(asctime)s [%(levelname)s] %(message)s"
logging.basicConfig(level=logging.INFO, format=LOG_FORMAT)
logger = logging.getLogger("collect_usda_gtr")


# ─────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────

def download_xlsx(url, save_path=None, timeout=60):
    logger.info(f"Downloading: {url}")
    resp = requests.get(url, timeout=timeout, verify=False)
    resp.raise_for_status()
    if save_path:
        Path(save_path).parent.mkdir(parents=True, exist_ok=True)
        with open(save_path, "wb") as f:
            f.write(resp.content)
    logger.info(f"  {len(resp.content):,} bytes")
    return resp.content


def open_sheet(content, sheet_name):
    """Open xlsx with read_only=False (required for correct parsing)."""
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def _j(val):
    """Convert to JSON-safe type."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip()
    if s in ("", "-", "  ", "n/a", "N/A"):
        return None
    try:
        return float(s.replace(",", ""))
    except ValueError:
        return s


# ─────────────────────────────────────────────
# PARSERS
# ─────────────────────────────────────────────

def parse_table1(content):
    """
    Sheet "Data" structure:
      Row 6: Date | Price | Rail | River | Gulf | PNW  (headers)
      Row 7+: weekly data from 2002-08-21
    """
    rows = open_sheet(content, "Data")

    records = []
    for row in rows[7:]:
        if len(row) < 6:
            continue
        if not isinstance(row[0], datetime):
            continue
        rec = {
            "date": row[0].strftime("%Y-%m-%d"),
            "diesel_usd_gal": _j(row[1]),
            "rail_usd_car": _j(row[2]),
            "barge_pct_tariff": _j(row[3]),
            "ocean_gulf_usd_mt": _j(row[4]),
            "ocean_pnw_usd_mt": _j(row[5]),
        }
        # Skip rows where everything is None
        if any(v is not None for k, v in rec.items() if k != "date"):
            records.append(rec)

    logger.info(f"  Table 1: {len(records)} weekly records" +
                (f" ({records[0]['date']} to {records[-1]['date']})" if records else ""))
    return records


def parse_table2(content):
    """
    Sheet "Data" structure:
      Row 1: _ | Commodity | Origin--destination | Origin Price | Destination Price | Price spreads
      Row 2+: data grouped in blocks of 5 per date
        Corn IL-Gulf, Corn NE-Gulf, Soybean IA-Gulf, HRW KS-Gulf, HRS ND-Portland
    """
    rows = open_sheet(content, "Data")

    records = []
    current_date = None

    for row in rows[2:]:
        if len(row) < 6:
            continue

        if isinstance(row[0], datetime):
            current_date = row[0].strftime("%Y-%m-%d")

        if current_date is None:
            continue

        commodity = str(row[1]).strip() if row[1] else None
        route = str(row[2]).strip() if row[2] else None
        if not commodity or not route:
            continue

        records.append({
            "date": current_date,
            "commodity": commodity,
            "route": route,
            "origin_price": _j(row[3]),
            "dest_price": _j(row[4]),
            "spread": _j(row[5]),
        })

    logger.info(f"  Table 2: {len(records)} spread records")

    # Latest by route
    latest_by_route = {}
    for rec in reversed(records):
        key = f"{rec['commodity']}|{rec['route']}"
        if key not in latest_by_route:
            latest_by_route[key] = rec

    return {"records": records, "latest_by_route": latest_by_route}


def parse_table7(content):
    """
    Sheet "GTR Table 7" structure:
      Row 0: _ | title with month reference
      Row 1: _ | Commodity | Railroad | Origin | Destination | Car Ownership | Tariff | Fuel surcharge
      Row 2+: current month snapshot (~40 routes)
    """
    rows = open_sheet(content, "GTR Table 7")

    title = str(rows[0][1]) if rows[0] and len(rows[0]) > 1 and rows[0][1] else ""

    records = []
    current_commodity = None

    for row in rows[2:]:
        if len(row) < 8:
            continue
        if all(row[i] is None for i in range(1, 8)):
            continue

        if row[1] and str(row[1]).strip():
            current_commodity = str(row[1]).strip()

        railroad = row[2]
        origin = row[3]
        if not railroad or not origin:
            continue

        records.append({
            "commodity": current_commodity,
            "railroad": str(railroad).strip(),
            "origin": str(origin).strip(),
            "destination": str(row[4]).strip() if row[4] else None,
            "car_ownership": str(row[5]).strip() if row[5] else None,
            "tariff_per_car": _j(row[6]),
            "fuel_surcharge_per_car": _j(row[7]),
        })

    logger.info(f"  Table 7: {len(records)} tariff records | {title}")
    return {"title": title, "records": records}


def parse_table9(content):
    """
    Sheet "Socrata_Data" structure:
      Row 0: headers (Week Ending, Corn_Lock 27, Corn_Lock 52, Corn_Lock 1, Corn_Total, ...)
      Row 1+: weekly data from 2003
    """
    rows = open_sheet(content, "Socrata_Data")

    headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]

    records = []
    for row in rows[1:]:
        if not row[0] or not isinstance(row[0], datetime):
            continue
        rec = {"date": row[0].strftime("%Y-%m-%d")}
        for i, h in enumerate(headers[1:], 1):
            if i < len(row):
                rec[h] = _j(row[i])
        records.append(rec)

    logger.info(f"  Table 9: {len(records)} weekly barge movement records")

    latest_corn = []
    for rec in records[-8:]:
        latest_corn.append({
            "date": rec["date"],
            "corn_total_tons": rec.get("Corn_Total"),
            "corn_4wk_avg": rec.get("Corn_4_weeks_AVG"),
        })

    return {"records": records, "latest_corn": latest_corn, "headers": headers}


# ─────────────────────────────────────────────
# SUMMARY
# ─────────────────────────────────────────────

def build_summary(parsed):
    summary = {}

    t1 = parsed.get("table1", [])
    if t1:
        for rec in reversed(t1):
            if any(v is not None for k, v in rec.items() if k != "date"):
                summary["cost_indicators"] = rec
                break

    t2 = parsed.get("table2", {}).get("latest_by_route", {})
    summary["spreads"] = {k: v for k, v in t2.items() if "Corn" in k or "Soy" in k}

    t7 = parsed.get("table7", {})
    summary["rail_tariffs"] = {"title": t7.get("title", ""), "route_count": len(t7.get("records", []))}

    return summary


# ─────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────

def save_json(data, filepath):
    Path(filepath).parent.mkdir(parents=True, exist_ok=True)
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False, default=str)
    logger.info(f"Saved: {filepath}")


def load_cache():
    f = CACHE_DIR / "usda_gtr_latest.json"
    if f.exists():
        try:
            return json.load(open(f, encoding="utf-8"))
        except Exception:
            pass
    return {}


def collect(tables=None, save_raw=True):
    now = datetime.now()
    ts = now.isoformat()

    if tables is None:
        tables = list(TABLES.keys())

    parsers = {"table1": parse_table1, "table2": parse_table2,
               "table7": parse_table7, "table9": parse_table9}

    parsed = {}
    errors = {}

    for key in tables:
        if key not in TABLES:
            continue
        info = TABLES[key]
        logger.info(f"Processing {key}: {info['label']}")
        try:
            raw_path = RAW_DIR / f"{key}.xlsx" if save_raw else None
            content = download_xlsx(info["url"], save_path=raw_path)
            parsed[key] = parsers[key](content)
        except Exception as e:
            logger.error(f"  Error on {key}: {e}")
            errors[key] = str(e)

    if not parsed:
        cached = load_cache()
        if cached:
            cached["status"] = "cached"
            return cached
        return {"source": "usda_gtr", "status": "error", "errors": errors}

    summary = build_summary(parsed)

    serialized = {}
    for key, data in parsed.items():
        if isinstance(data, list):
            serialized[key] = {"label": TABLES[key]["label"], "count": len(data), "records": data}
        else:
            serialized[key] = {"label": TABLES[key]["label"], **data}
            if "records" in data:
                serialized[key]["count"] = len(data["records"])

    output = {
        "source": "usda_gtr",
        "source_url": "https://www.ams.usda.gov/services/transportation-analysis/gtr-datasets",
        "collection_timestamp": ts,
        "status": "ok",
        "errors": errors if errors else None,
        "summary": summary,
        "data": serialized,
    }

    timestamp = now.strftime("%Y%m%d_%H%M%S")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    save_json(output, OUTPUT_DIR / f"usda_gtr_{timestamp}.json")
    save_json(output, OUTPUT_DIR / "usda_gtr_latest.json")
    save_json(output, CACHE_DIR / "usda_gtr_latest.json")

    # Log
    logger.info("=" * 60)
    logger.info("USDA GTR - COLLECTION SUMMARY")
    logger.info("=" * 60)
    if "cost_indicators" in summary:
        ci = summary["cost_indicators"]
        logger.info(f"  Latest cost indicators ({ci['date']}):")
        logger.info(f"    Diesel:       ${ci.get('diesel_usd_gal', '?')}/gal")
        logger.info(f"    Rail:         ${ci.get('rail_usd_car', '?')}/car")
        logger.info(f"    Barge:        {ci.get('barge_pct_tariff', '?')}% of tariff")
        logger.info(f"    Ocean Gulf:   ${ci.get('ocean_gulf_usd_mt', '?')}/mt")
        logger.info(f"    Ocean PNW:    ${ci.get('ocean_pnw_usd_mt', '?')}/mt")
    if summary.get("spreads"):
        logger.info("  Latest price spreads:")
        for route, data in summary["spreads"].items():
            logger.info(f"    {data['commodity']} {data['route']}: ${data['spread']}/bu")
    rt = summary.get("rail_tariffs", {})
    logger.info(f"  Rail tariffs: {rt.get('route_count', 0)} routes | {rt.get('title', '')}")
    for key in serialized:
        logger.info(f"  {TABLES[key]['label']:.<45} {serialized[key].get('count', '?'):>5} records")
    logger.info("=" * 60)

    return output


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="AgriMacro - USDA GTR Collector")
    parser.add_argument("--tables", nargs="+", default=None,
                        choices=["table1", "table2", "table7", "table9"])
    parser.add_argument("--output-dir", type=str, default=None)
    args = parser.parse_args()

    if args.output_dir:
        OUTPUT_DIR = Path(args.output_dir)
        CACHE_DIR = OUTPUT_DIR / "cache"
        RAW_DIR = OUTPUT_DIR / "raw_xlsx"

    result = collect(tables=args.tables)
    sys.exit(0 if result["status"] in ("ok", "cached") else 1)
